from datetime import date
from pathlib import Path

import pytest
from openpyxl import Workbook

from compass import rvtools
from compass.__main__ import main
from compass.analyze import analyze
from compass.cloud import cloud_cost, load_catalog, pick_instance
from compass.inventory import summarize
from compass.licensing import renewal_cost
from compass.onprem import hyperv_cost, nutanix_cost, olvm_cost
from compass.proxmox import proxmox_cost
from compass.rules import load_rules
from compass.rvtools import VM, Host, RVToolsError

FIX = Path(__file__).parent / "fixtures"
TODAY = date(2026, 10, 28)


@pytest.fixture(scope="module")
def broadcom():
    return load_rules("broadcom")


@pytest.fixture(scope="module")
def inv():
    return rvtools.load(FIX / "sample_a.xlsx")


def host(sockets, cpc, name="h"):
    return Host(name, "c", "d", sockets, cpc, None, "", "")


def vm(name="v", vcpus=2, mem_gib=8, os="Ubuntu Linux (64-bit)", on=True, disk_gib=100):
    return VM(name, on, False, vcpus, mem_gib * 1024, None, None, os, "", "", "", [disk_gib * 1024])


# --- lector ---

def test_both_layouts_read_the_same(inv):
    other = rvtools.load(FIX / "sample_b.xlsx")
    key = lambda i: [(v.name, v.powered_on, v.vcpus, v.memory_mib, v.disk_total_mib, v.template) for v in i.vms]
    assert key(inv) == key(other)
    assert [(h.name, h.sockets, h.cores_per_cpu) for h in inv.hosts] == \
           [(h.name, h.sockets, h.cores_per_cpu) for h in other.hosts]


def test_missing_sheet(tmp_path):
    wb = Workbook()
    wb.active.title = "vInfo"
    wb.save(tmp_path / "x.xlsx")
    with pytest.raises(RVToolsError, match="vHost"):
        rvtools.load(tmp_path / "x.xlsx")


def test_missing_column_names_alternatives(tmp_path):
    wb = Workbook()
    wb.active.title = "vInfo"
    wb.active.append(["VM", "Powerstate", "CPUs"])
    for s in ("vHost", "vDisk"):
        wb.create_sheet(s).append(["Host", "# CPU"] if s == "vHost" else ["VM", "Capacity MiB"])
    wb.save(tmp_path / "x.xlsx")
    with pytest.raises(RVToolsError, match='"Memory"'):
        rvtools.load(tmp_path / "x.xlsx")


# --- inventario ---

def test_summary(inv):
    s = summarize(inv, load_rules("os_eol"), TODAY)
    assert s["vms_total"] == 39 and s["templates"] == 1
    assert s["vms_powered_off"] == 5  # vm-000, 009, 018, 027, 036
    assert s["hosts"] == 4 and s["sockets"] == 7 and s["physical_cores"] == 96
    labels = {u["label"]: u["state"] for u in s["unsupported_os"]}
    assert labels["Windows Server 2012 / 2012 R2"] == "expired"
    assert labels["CentOS 7"] == "expired"
    assert labels["Windows Server 2016"] == "expiring"
    assert "RHEL 8" not in labels


# --- licenciamiento ---

def test_cores_padded_to_16_per_cpu(inv, broadcom):
    r = renewal_cost(inv.hosts, broadcom, "VCF")
    # 2x12 -> 32, 2x12 -> 32, 2x20 -> 40, 1x8 -> 16
    assert r["licensed_cores"] == 120 and r["physical_cores"] == 96 and r["padding_cores"] == 24
    assert r["annual_list_usd"] == {"low": 120 * 350, "high": 120 * 400}
    assert r["vsan_tib_included"] == 120


def test_vvf_vsan_rounds_up(broadcom):
    r = renewal_cost([host(1, 18)], broadcom, "VVF")
    assert r["vsan_tib_included"] == 5  # 18 * 0.25 = 4.5 -> 5


def test_order_minimum_is_opt_in(broadcom):
    assert renewal_cost([host(1, 8)], broadcom, "VCF")["licensed_cores"] == 16
    assert renewal_cost([host(1, 8)], broadcom, "VCF", apply_order_minimum=True)["licensed_cores"] == 72


def test_late_surcharge_only_first_year(broadcom):
    r = renewal_cost([host(2, 16)], broadcom, "VCF", term_years=3, late_renewal=True)
    annual = 32 * 350
    assert r["term_list_usd"]["low"] == round(annual * 3 + annual * 0.2)


def test_prepared_negotiation_is_cheaper(broadcom):
    n = renewal_cost([host(2, 16)], broadcom, "VCF")["annual_negotiated_usd"]
    assert n["prepared"]["high"] < n["unprepared"]["high"]
    assert n["prepared"]["low"] == round(32 * 350 * (1 - 0.55))


# --- nube ---

def test_picks_cheapest_fitting_instance():
    cat = load_catalog(FIX / "catalog_sample.json")
    assert pick_instance(vm(vcpus=4, mem_gib=8), cat)["name"] == "c-4-8"
    assert pick_instance(vm(vcpus=3, mem_gib=20), cat)["name"] == "m-4-32"
    assert pick_instance(vm(vcpus=64, mem_gib=8), cat) is None


def test_cloud_cost_ranges_and_skips_off():
    cat = load_catalog(FIX / "catalog_sample.json")
    vms = [vm("a"), vm("b", os="Microsoft Windows Server 2022 (64-bit)"), vm("c", on=False), vm("big", vcpus=64)]
    r = cloud_cost(vms, cat)
    assert r["vms_placed"] == 2 and r["vms_unplaced"] == ["big"]
    compute = (0.1 + 0.2) * 730
    storage = 2 * 100 * 0.1
    assert r["annual_usd"]["high"] == round((compute + storage) * 12)
    assert r["annual_usd"]["low"] == round((compute * 0.5 + storage) * 12)
    assert r["catalog_is_sample"] is True


# --- Proxmox ---

def test_proxmox_per_socket(inv):
    r = proxmox_cost(inv.hosts, load_rules("proxmox"))
    assert r["sockets"] == 7 and r["annual_eur"] == 7 * 550
    assert r["upper_tier"] == "premium"


# --- Hyper-V, Nutanix, Oracle ---

def test_hyperv_core_minimums(inv):
    r = hyperv_cost(inv.hosts, load_rules("hyperv"), term_years=3)
    # 24 + 24 + 40 + (1x8 sube a 16 por servidor)
    assert r["licensed_cores"] == 104
    license_usd = 104 * 6771 / 16
    assert r["annual_usd"] == {"low": round(license_usd / 3), "high": round(license_usd / 3 + license_usd * 0.25)}


def test_hyperv_existing_datacenter_pays_only_sa(inv):
    r = hyperv_cost(inv.hosts, load_rules("hyperv"), owns_datacenter=True)
    assert r["annual_usd"] == {"low": 0, "high": round(104 * 6771 / 16 * 0.25)}


def test_hyperv_min_8_per_cpu():
    assert hyperv_cost([host(2, 6)], load_rules("hyperv"))["licensed_cores"] == 16
    assert hyperv_cost([host(4, 6)], load_rules("hyperv"))["licensed_cores"] == 32


def test_nutanix_physical_cores_with_discount(inv):
    r = nutanix_cost(inv.hosts, load_rules("nutanix"))
    assert r["licensed_cores"] == 96
    assert r["annual_usd"] == {"low": round(96 * 649 * 0.45), "high": round(96 * 649 * 0.65)}


def test_olvm_rounds_cpus_to_pairs(inv):
    r = olvm_cost(inv.hosts, load_rules("oracle"))
    assert r["cpu_pairs"] == 4  # el host de 1 CPU cuenta como un par
    assert r["annual_usd"] == {"low": 4 * 1399, "high": 4 * 2499}
    assert olvm_cost([host(3, 8)], load_rules("oracle"))["cpu_pairs"] == 2


# --- de punta a punta ---

def test_analyze_and_cli(capsys):
    r = analyze(FIX / "sample_b.xlsx", [FIX / "catalog_sample.json"], today=TODAY)
    assert set(r["scenarios"]) == {"VCF", "VVF", "SAMPLE", "PROXMOX", "HYPERV", "NUTANIX", "OLVM"}
    for key in ("HYPERV", "NUTANIX", "OLVM"):
        assert "region" not in r["scenarios"][key] and r["scenarios"][key]["basis"]
    assert main([str(FIX / "sample_a.xlsx"), "--catalog", str(FIX / "catalog_sample.json")]) == 0
    out = capsys.readouterr().out
    assert "PRECIOS DE MUESTRA" in out and "120 núcleos licenciados" in out


# --- reporte ---

def test_report_html_marks_sample_prices_and_lists_sources():
    from compass.report import render_html
    r = analyze(FIX / "sample_a.xlsx", [FIX / "catalog_sample.json"], today=TODAY)
    html = render_html(r, client="Cliente <prueba>", today=TODAY)
    assert "Datos de ejemplo" in html  # el catálogo de muestra fuerza la marca
    assert "Cliente &lt;prueba&gt;" in html
    assert "KB 313548" in html and "En disputa" in html
    assert "Renovar VMware Cloud Foundation" in html and "Migrar a Proxmox VE" in html
