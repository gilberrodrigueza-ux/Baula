"""Genera exports sintéticos de RVTools para pruebas. No son datos reales.

Uso: python scripts/make_sample_rvtools.py tests/fixtures

Crea dos archivos con el mismo inventario y distinta forma:
- sample_a.xlsx: encabezados en MiB y el orden habitual.
- sample_b.xlsx: columnas en otro orden, encabezados en MB y columnas extra,
  para probar que el lector busca por nombre. Las diferencias reales entre
  RVTools 4.7 y 4.8 se confirman con los exports de la fase 0.
"""
from __future__ import annotations

import random
import sys
from pathlib import Path

from openpyxl import Workbook

OSES = [
    "Microsoft Windows Server 2019 (64-bit)",
    "Microsoft Windows Server 2022 (64-bit)",
    "Microsoft Windows Server 2012 R2 (64-bit)",
    "Microsoft Windows Server 2016 (64-bit)",
    "Red Hat Enterprise Linux 8 (64-bit)",
    "CentOS 7 (64-bit)",
    "Ubuntu Linux (64-bit)",
]
HOSTS = [  # nombre, cluster, sockets, núcleos por CPU, RAM GiB
    ("esx01.lab.local", "PROD", 2, 12, 512),
    ("esx02.lab.local", "PROD", 2, 12, 512),
    ("esx03.lab.local", "PROD", 2, 20, 768),
    ("esx04.lab.local", "DR", 1, 8, 256),
]


def inventory(seed: int = 7):
    rnd = random.Random(seed)
    vms = []
    for i in range(40):
        host = HOSTS[i % len(HOSTS)]
        cpus = rnd.choice([1, 2, 2, 4, 4, 8, 16])
        mem_gib = rnd.choice([2, 4, 8, 8, 16, 32, 64])
        disks = [rnd.choice([40, 60, 100]) * 1024] + ([rnd.choice([100, 250, 500]) * 1024] if rnd.random() < 0.5 else [])
        vms.append({
            "VM": f"vm-{i:03d}",
            "Powerstate": "poweredOff" if i % 9 == 0 else "poweredOn",
            "Template": i == 39,
            "CPUs": cpus,
            "Memory": mem_gib * 1024,
            "Provisioned": sum(disks) + mem_gib * 1024,
            "In Use": int(sum(disks) * 0.55),
            "OS": OSES[i % len(OSES)],
            "Host": host[0],
            "Cluster": host[1],
            "Datacenter": "DC-SJO",
            "disks": disks,
        })
    return vms


def write(path: Path, variant: str):
    vms = inventory()
    wb = Workbook()
    ws = wb.active
    ws.title = "vInfo"
    if variant == "a":
        cols = ["VM", "Powerstate", "Template", "CPUs", "Memory", "Provisioned MiB", "In Use MiB",
                "OS according to the configuration file", "OS according to the VMware Tools", "Host", "Cluster", "Datacenter"]
    else:
        cols = ["Datacenter", "Cluster", "Host", "VM", "Annotation", "OS according to the VMware Tools",
                "Template", "Powerstate", "Memory MB", "CPUs", "In Use MB", "Provisioned MB",
                "OS according to the configuration file"]
    ws.append(cols)
    for v in vms:
        val = {"Provisioned MiB": v["Provisioned"], "Provisioned MB": v["Provisioned"],
               "In Use MiB": v["In Use"], "In Use MB": v["In Use"],
               "Memory MB": v["Memory"], "Annotation": "",
               "OS according to the configuration file": v["OS"],
               "OS according to the VMware Tools": v["OS"] if v["Powerstate"] == "poweredOn" else "",
               **v}
        ws.append([val[c] for c in cols])

    wh = wb.create_sheet("vHost")
    hcols = ["Host", "Datacenter", "Cluster", "CPU Model", "# CPU", "Cores per CPU", "# Cores", "# Memory", "ESX Version"]
    if variant == "b":
        hcols = ["Datacenter", "Cluster", "Host", "# Cores", "# CPU", "# Memory", "ESX Version", "CPU Model"]
    wh.append(hcols)
    for name, cluster, sockets, cpc, ram in HOSTS:
        val = {"Host": name, "Datacenter": "DC-SJO", "Cluster": cluster, "CPU Model": "Intel(R) Xeon(R) Gold",
               "# CPU": sockets, "Cores per CPU": cpc, "# Cores": sockets * cpc, "# Memory": ram * 1024,
               "ESX Version": "VMware ESXi 8.0.3"}
        wh.append([val[c] for c in hcols])

    wd = wb.create_sheet("vDisk")
    dcol = "Capacity MiB" if variant == "a" else "Capacity MB"
    wd.append(["VM", "Disk", dcol, "Thin"])
    for v in vms:
        for n, cap in enumerate(v["disks"], 1):
            wd.append([v["VM"], f"Hard disk {n}", cap, True])
    wb.save(path)


if __name__ == "__main__":
    out = Path(sys.argv[1] if len(sys.argv) > 1 else "tests/fixtures")
    out.mkdir(parents=True, exist_ok=True)
    write(out / "sample_a.xlsx", "a")
    write(out / "sample_b.xlsx", "b")
    print(f"Exports sintéticos en {out}")
