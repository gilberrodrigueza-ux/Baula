"""Genera el export de RVTools de la empresa ficticia del reporte de ejemplo. No son datos reales.

Uso: python scripts/make_demo_rvtools.py examples/demo/empresa-ejemplo.xlsx

"Empresa Ejemplo S.A." es una mediana de servicios en San José con dos sitios:
6 hosts de producción y 2 de DR, unas 180 VMs, algo de Windows Server 2012 R2 y CentOS 7
heredado, y VMs apagadas que nadie reclama. Sirve para mostrar el reporte en entrevistas.
"""
from __future__ import annotations

import random
import sys
from pathlib import Path

from openpyxl import Workbook

OSES = [  # SO, peso
    ("Microsoft Windows Server 2019 (64-bit)", 30),
    ("Microsoft Windows Server 2022 (64-bit)", 22),
    ("Microsoft Windows Server 2016 (64-bit)", 14),
    ("Microsoft Windows Server 2012 R2 (64-bit)", 9),
    ("Red Hat Enterprise Linux 8 (64-bit)", 12),
    ("Red Hat Enterprise Linux 9 (64-bit)", 6),
    ("CentOS 7 (64-bit)", 5),
    ("Ubuntu Linux (64-bit)", 8),
]
HOSTS = [  # nombre, cluster, datacenter, sockets, núcleos por CPU, RAM GiB, versión
    *[(f"esx-prd-0{i}.ejemplo.local", "PROD", "DC-SJO", 2, 24, 1024, "VMware ESXi 8.0.3") for i in range(1, 5)],
    *[(f"esx-prd-0{i}.ejemplo.local", "PROD-LEGACY", "DC-SJO", 2, 12, 512, "VMware ESXi 8.0.2") for i in range(5, 7)],
    *[(f"esx-dr-0{i}.ejemplo.local", "DR", "DC-HER", 2, 10, 384, "VMware ESXi 8.0.2") for i in range(1, 3)],
]
ROLES = ["app", "db", "web", "file", "ad", "erp", "bi", "mon", "test", "dev"]


def inventory(seed: int = 2026):
    rnd = random.Random(seed)
    oses = [o for o, w in OSES for _ in range(w)]
    vms = []
    for i in range(184):
        role = ROLES[i % len(ROLES)]
        host = HOSTS[i % 6] if i % 8 else HOSTS[6 + (i // 8) % 2]
        cpus = rnd.choice([2, 2, 2, 4, 4, 4, 8, 8, 16]) if role != "db" else rnd.choice([8, 8, 16])
        mem_gib = rnd.choice([4, 8, 8, 16, 16, 32]) if role != "db" else rnd.choice([32, 64, 64])
        disks = [rnd.choice([60, 80, 120]) * 1024]
        if role in ("db", "file", "erp", "bi") or rnd.random() < 0.3:
            disks.append(rnd.choice([100, 250, 500, 1000]) * 1024)
        powered_off = role in ("test", "dev") and rnd.random() < 0.55 or rnd.random() < 0.04
        vms.append({
            "VM": f"{role}-{i:03d}",
            "Powerstate": "poweredOff" if powered_off else "poweredOn",
            "Template": i in (182, 183),
            "CPUs": cpus,
            "Memory": mem_gib * 1024,
            "Provisioned MiB": sum(disks) + mem_gib * 1024,
            "In Use MiB": int(sum(disks) * rnd.uniform(0.35, 0.8)),
            "OS": rnd.choice(oses),
            "Host": host[0],
            "Cluster": host[1],
            "Datacenter": host[2],
            "disks": disks,
        })
    return vms


def write(path: Path):
    vms = inventory()
    wb = Workbook()
    ws = wb.active
    ws.title = "vInfo"
    cols = ["VM", "Powerstate", "Template", "CPUs", "Memory", "Provisioned MiB", "In Use MiB",
            "OS according to the configuration file", "OS according to the VMware Tools", "Host", "Cluster", "Datacenter"]
    ws.append(cols)
    for v in vms:
        val = {**v, "OS according to the configuration file": v["OS"],
               "OS according to the VMware Tools": v["OS"] if v["Powerstate"] == "poweredOn" else ""}
        ws.append([val[c] for c in cols])

    wh = wb.create_sheet("vHost")
    wh.append(["Host", "Datacenter", "Cluster", "CPU Model", "# CPU", "Cores per CPU", "# Cores", "# Memory", "ESX Version"])
    for name, cluster, dc, sockets, cpc, ram, esx in HOSTS:
        wh.append([name, dc, cluster, "Intel(R) Xeon(R) Gold", sockets, cpc, sockets * cpc, ram * 1024, esx])

    wd = wb.create_sheet("vDisk")
    wd.append(["VM", "Disk", "Capacity MiB", "Thin"])
    for v in vms:
        for n, cap in enumerate(v["disks"], 1):
            wd.append([v["VM"], f"Hard disk {n}", cap, True])
    wb.save(path)


if __name__ == "__main__":
    out = Path(sys.argv[1] if len(sys.argv) > 1 else "examples/demo/empresa-ejemplo.xlsx")
    out.parent.mkdir(parents=True, exist_ok=True)
    write(out)
    print(f"Export ficticio en {out}")
