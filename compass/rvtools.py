"""Lector de exports de RVTools (.xlsx).

Busca las hojas y las columnas por nombre, nunca por posición, porque el orden
cambia entre versiones de RVTools. Cada campo acepta varios encabezados
posibles (por ejemplo "Memory" o "Memory MiB") y se normalizan a MiB.
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook

REQUIRED_SHEETS = ("vInfo", "vHost", "vDisk")

# campo interno -> encabezados aceptados, en orden de preferencia.
# Un sufijo " MB" se trata como MiB: RVTools usó MB y MiB para el mismo valor.
VINFO_COLUMNS = {
    "name": ["VM"],
    "powerstate": ["Powerstate", "Power state"],
    "template": ["Template"],
    "cpus": ["CPUs", "Num CPU"],
    "memory_mib": ["Memory", "Memory MiB", "Memory MB"],
    "provisioned_mib": ["Provisioned MiB", "Provisioned MB"],
    "in_use_mib": ["In Use MiB", "In Use MB"],
    "os_config": ["OS according to the configuration file", "OS"],
    "os_tools": ["OS according to the VMware Tools"],
    "host": ["Host"],
    "cluster": ["Cluster"],
    "datacenter": ["Datacenter"],
}
VINFO_REQUIRED = ("name", "powerstate", "cpus", "memory_mib")

VHOST_COLUMNS = {
    "name": ["Host"],
    "cluster": ["Cluster"],
    "datacenter": ["Datacenter"],
    "sockets": ["# CPU", "CPU sockets"],
    "cores_per_cpu": ["Cores per CPU"],
    "cores": ["# Cores"],
    "memory_mib": ["# Memory", "Memory MiB", "Memory MB"],
    "cpu_model": ["CPU Model"],
    "esx_version": ["ESX Version"],
}
VHOST_REQUIRED = ("name", "sockets")

VDISK_COLUMNS = {
    "vm": ["VM"],
    "disk": ["Disk"],
    "capacity_mib": ["Capacity MiB", "Capacity MB"],
    "thin": ["Thin"],
}
VDISK_REQUIRED = ("vm", "capacity_mib")


class RVToolsError(ValueError):
    """El archivo no tiene la forma de un export de RVTools usable."""


@dataclass
class VM:
    name: str
    powered_on: bool
    template: bool
    vcpus: int
    memory_mib: float
    provisioned_mib: float | None
    in_use_mib: float | None
    os: str
    host: str
    cluster: str
    datacenter: str
    disks_mib: list[float] = field(default_factory=list)

    @property
    def disk_total_mib(self) -> float:
        if self.disks_mib:
            return sum(self.disks_mib)
        return self.provisioned_mib or 0.0


@dataclass
class Host:
    name: str
    cluster: str
    datacenter: str
    sockets: int
    cores_per_cpu: int
    memory_mib: float | None
    cpu_model: str
    esx_version: str


@dataclass
class Inventory:
    source: str
    vms: list[VM]
    hosts: list[Host]
    warnings: list[str] = field(default_factory=list)


def _norm(header: object) -> str:
    return re.sub(r"\s+", " ", str(header or "")).strip().lower()


def _map_columns(headers: list[object], spec: dict[str, list[str]], required: tuple[str, ...], sheet: str) -> dict[str, int]:
    index = {_norm(h): i for i, h in enumerate(headers) if h is not None}
    found = {}
    for key, names in spec.items():
        for name in names:
            if _norm(name) in index:
                found[key] = index[_norm(name)]
                break
    missing = [k for k in required if k not in found]
    if missing:
        wanted = ", ".join(" o ".join(f'"{n}"' for n in spec[k]) for k in missing)
        raise RVToolsError(f"La hoja {sheet} no tiene las columnas {wanted}.")
    return found


def _rows(ws, spec, required, sheet):
    it = ws.iter_rows(values_only=True)
    try:
        headers = list(next(it))
    except StopIteration:
        raise RVToolsError(f"La hoja {sheet} está vacía.")
    cols = _map_columns(headers, spec, required, sheet)
    for row in it:
        if row is None or all(v is None for v in row):
            continue
        yield {k: (row[i] if i < len(row) else None) for k, i in cols.items()}


def _num(value, default=None):
    if value is None or value == "":
        return default
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return float(str(value).replace(",", ""))
    except ValueError:
        return default


def _bool(value) -> bool:
    if isinstance(value, bool):
        return value
    return str(value).strip().lower() in ("true", "yes", "1", "sí", "si")


def _str(value) -> str:
    return "" if value is None else str(value).strip()


def load(path: str | Path) -> Inventory:
    path = Path(path)
    wb = load_workbook(path, read_only=True, data_only=True)
    sheets = {name.lower(): name for name in wb.sheetnames}
    missing = [s for s in REQUIRED_SHEETS if s.lower() not in sheets]
    if missing:
        raise RVToolsError(
            f"Faltan las hojas {', '.join(missing)}. Exportá desde RVTools con 'Export all to xlsx'."
        )
    warnings: list[str] = []

    vms: dict[str, VM] = {}
    for r in _rows(wb[sheets["vinfo"]], VINFO_COLUMNS, VINFO_REQUIRED, "vInfo"):
        name = _str(r["name"])
        if not name:
            continue
        os_name = _str(r.get("os_tools")) or _str(r.get("os_config"))
        vm = VM(
            name=name,
            powered_on=_str(r["powerstate"]).lower() in ("poweredon", "powered on"),
            template=_bool(r.get("template")),
            vcpus=int(_num(r["cpus"], 0)),
            memory_mib=_num(r["memory_mib"], 0.0),
            provisioned_mib=_num(r.get("provisioned_mib")),
            in_use_mib=_num(r.get("in_use_mib")),
            os=os_name,
            host=_str(r.get("host")),
            cluster=_str(r.get("cluster")),
            datacenter=_str(r.get("datacenter")),
        )
        if name in vms:
            warnings.append(f"VM duplicada en vInfo: {name}; se usa la primera.")
            continue
        vms[name] = vm

    for r in _rows(wb[sheets["vdisk"]], VDISK_COLUMNS, VDISK_REQUIRED, "vDisk"):
        vm = vms.get(_str(r["vm"]))
        cap = _num(r["capacity_mib"])
        if vm is not None and cap is not None:
            vm.disks_mib.append(cap)

    hosts: list[Host] = []
    for r in _rows(wb[sheets["vhost"]], VHOST_COLUMNS, VHOST_REQUIRED, "vHost"):
        name = _str(r["name"])
        if not name:
            continue
        sockets = int(_num(r["sockets"], 0))
        cpc = _num(r.get("cores_per_cpu"))
        if cpc is None:
            total = _num(r.get("cores"))
            if total is None or sockets == 0:
                raise RVToolsError(f"El host {name} no tiene 'Cores per CPU' ni '# Cores'.")
            cpc = total / sockets
        hosts.append(Host(
            name=name,
            cluster=_str(r.get("cluster")),
            datacenter=_str(r.get("datacenter")),
            sockets=sockets,
            cores_per_cpu=int(cpc),
            memory_mib=_num(r.get("memory_mib")),
            cpu_model=_str(r.get("cpu_model")),
            esx_version=_str(r.get("esx_version")),
        ))

    if not hosts:
        warnings.append("La hoja vHost no tiene hosts: no se puede calcular licenciamiento por núcleo.")
    return Inventory(source=path.name, vms=list(vms.values()), hosts=hosts, warnings=warnings)
